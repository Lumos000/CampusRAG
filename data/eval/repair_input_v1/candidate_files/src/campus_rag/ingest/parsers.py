from __future__ import annotations

import hashlib
import re
from datetime import datetime, timezone
from pathlib import Path

from campus_rag.types import ParsedDocument, ParsedTable, SourceDocument
from campus_rag.ingest.navigation import navigation_values


def parse_file(path: str | Path, source_url: str = "", department: str = "", published_at: str | None = None) -> ParsedDocument:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    source = SourceDocument(
        title=file_path.stem,
        source_url=source_url or file_path.resolve().as_uri(),
        department=department,
        published_at=published_at,
        crawled_at=datetime.now(timezone.utc).isoformat(),
        local_path=str(file_path),
        content_type=suffix.lstrip("."),
        sha256=hashlib.sha256(file_path.read_bytes()).hexdigest(),
    )
    if suffix in {".html", ".htm"}:
        return _parse_html(file_path, source)
    if suffix == ".pdf":
        return _parse_pdf(file_path, source)
    if suffix == ".docx":
        return _parse_docx(file_path, source)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx(file_path, source)
    if suffix in {".txt", ".md"}:
        return ParsedDocument(source, [(source.title, file_path.read_text(encoding="utf-8"))])
    raise ValueError(f"不支持的文件类型：{suffix}")


def _parse_html(path: Path, source: SourceDocument) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "lxml")
    if soup.title and soup.title.string:
        source.title = soup.title.string.strip()
    for node in soup(["script", "style", "nav", "footer", "noscript"]):
        node.decompose()
    article = soup.select_one('.wp_articlecontent, .Article_Content, .article-content, article')
    if article is not None:
        soup = article
    tables: list[ParsedTable] = []
    for index, table in enumerate(soup.find_all("table"), 1):
        grid = _html_table_grid(table)
        cells = table.select('td, th')
        fraction = sum(bool(cell.select_one('a[href]')) for cell in cells) / max(1, len(cells))
        if navigation_values([v for row in grid for v in row], fraction):
            table.decompose()
            continue
        if grid:
            headers, rows, levels = _choose_headers(grid)
            caption = table.find("caption")
            tables.append(ParsedTable(
                name=caption.get_text(" ", strip=True) if caption else f"表{index}",
                headers=headers, rows=rows, locator=f"html-table-{index}", header_levels=levels,
            ))
        table.decompose()
    sections = _html_sections(soup)
    return ParsedDocument(source, sections, tables)


def _html_table_grid(table) -> list[list[str]]:
    grid: list[list[str]] = []
    spans: dict[tuple[int, int], str] = {}
    for row_index, tr in enumerate(table.find_all("tr")):
        row: list[str] = []
        col = 0
        cells = tr.find_all(["th", "td"], recursive=False)
        for cell in cells:
            while (row_index, col) in spans:
                row.append(spans[(row_index, col)])
                col += 1
            value = cell.get_text(" ", strip=True)
            rowspan = int(cell.get("rowspan", 1) or 1)
            colspan = int(cell.get("colspan", 1) or 1)
            for offset in range(colspan):
                row.append(value)
                for future in range(1, rowspan):
                    spans[(row_index + future, col + offset)] = value
            col += colspan
        while (row_index, col) in spans:
            row.append(spans[(row_index, col)])
            col += 1
        if any(row):
            grid.append(row)
    return _rectangular(grid)


def _html_sections(soup) -> list[tuple[str, str]]:
    sections: list[tuple[str, str]] = []
    heading = "正文"
    buffer: list[str] = []
    for node in soup.find_all(["h1", "h2", "h3", "h4", "p", "li"]):
        text = node.get_text(" ", strip=True)
        if not text:
            continue
        if node.name.startswith("h"):
            if buffer:
                sections.append((heading, "\n\n".join(buffer)))
            heading, buffer = text, []
        else:
            buffer.append(text)
    if buffer:
        sections.append((heading, "\n\n".join(buffer)))
    if not sections:
        text = soup.get_text('\n', strip=True)
        if text:
            sections.append((heading, text))
    return sections


def _parse_pdf(path: Path, source: SourceDocument) -> ParsedDocument:
    import fitz

    sections: list[tuple[str, str]] = []
    tables: list[ParsedTable] = []
    warnings: list[str] = []
    with fitz.open(path) as document:
        for page_index, page in enumerate(document, 1):
            text = page.get_text("text").strip()
            if text:
                sections.append((f"第{page_index}页", text))
            elif page.get_images():
                warnings.append(f"第{page_index}页疑似扫描页，需要 OCR 或人工补录")
            if hasattr(page, "find_tables"):
                try:
                    found = page.find_tables()
                    for table_index, item in enumerate(found.tables, 1):
                        grid = [["" if value is None else str(value) for value in row] for row in item.extract()]
                        headers, rows, levels = _choose_headers(_rectangular(grid))
                        if headers:
                            tables.append(ParsedTable(
                                name=f"第{page_index}页表{table_index}", headers=headers, rows=rows,
                                locator=f"page-{page_index}-table-{table_index}", header_levels=levels,
                            ))
                except Exception as exc:  # parser failure must be visible in audit
                    warnings.append(f"第{page_index}页表格解析失败：{exc}")
    return ParsedDocument(source, sections, tables, warnings)


def _parse_docx(path: Path, source: SourceDocument) -> ParsedDocument:
    from docx import Document

    document = Document(path)
    sections: list[tuple[str, str]] = []
    heading = source.title
    buffer: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            if buffer:
                sections.append((heading, "\n\n".join(buffer)))
            heading, buffer = text, []
        else:
            buffer.append(text)
    if buffer:
        sections.append((heading, "\n\n".join(buffer)))
    tables: list[ParsedTable] = []
    for index, table in enumerate(document.tables, 1):
        grid = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        headers, rows, levels = _choose_headers(_rectangular(grid))
        if headers:
            tables.append(ParsedTable(f"表{index}", headers, rows, f"docx-table-{index}", levels))
    return ParsedDocument(source, sections, tables)


def _parse_xlsx(path: Path, source: SourceDocument) -> ParsedDocument:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=False)
    tables: list[ParsedTable] = []
    for sheet in workbook.worksheets:
        merged_values: dict[tuple[int, int], object] = {}
        for merged in sheet.merged_cells.ranges:
            value = sheet.cell(merged.min_row, merged.min_col).value
            for row in range(merged.min_row, merged.max_row + 1):
                for column in range(merged.min_col, merged.max_col + 1):
                    merged_values[(row, column)] = value
        grid: list[list[str]] = []
        for row in range(1, sheet.max_row + 1):
            values = []
            for column in range(1, sheet.max_column + 1):
                value = merged_values.get((row, column), sheet.cell(row, column).value)
                values.append("" if value is None else str(value).strip())
            if any(values):
                grid.append(values)
        if grid:
            headers, rows, levels = _choose_headers(_rectangular(grid))
            tables.append(ParsedTable(sheet.title, headers, rows, f"sheet:{sheet.title}", levels))
    workbook.close()
    return ParsedDocument(source, [], tables)


def _choose_headers(grid: list[list[str]]) -> tuple[list[str], list[list[str]], int]:
    if not grid:
        return [], [], 0
    first = grid[0]
    use_two = len(grid) > 1 and _looks_like_multilevel_header(first, grid[1])
    if use_two:
        headers = []
        for index, (upper, lower) in enumerate(zip(first, grid[1])):
            parts = [value for value in (upper.strip(), lower.strip()) if value]
            headers.append(" / ".join(dict.fromkeys(parts)) or f"列{index + 1}")
        rows, levels = grid[2:], 2
    else:
        headers = [value or f"列{index + 1}" for index, value in enumerate(first)]
        rows, levels = grid[1:], 1
    seen: dict[str, int] = {}
    unique: list[str] = []
    for header in headers:
        seen[header] = seen.get(header, 0) + 1
        unique.append(header if seen[header] == 1 else f"{header}_{seen[header]}")
    return unique, rows, levels


def _looks_like_multilevel_header(first: list[str], second: list[str]) -> bool:
    nonempty = [value for value in first if value]
    repeated = len(set(nonempty)) < len(nonempty)
    blanks = sum(not value for value in first) >= max(1, len(first) // 3)
    second_textual = sum(not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", value.replace(",", "")) for value in second if value)
    return (repeated or blanks) and second_textual >= max(1, len([value for value in second if value]) // 2)


def _rectangular(grid: list[list[str]]) -> list[list[str]]:
    width = max((len(row) for row in grid), default=0)
    return [row + [""] * (width - len(row)) for row in grid]
