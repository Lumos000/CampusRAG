from openpyxl import Workbook

from campus_rag.ingest.parsers import parse_file


def test_xlsx_merged_cells_are_expanded(tmp_path):
    path = tmp_path / "table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "招生计划"
    sheet.append(["专业", "选科", "计划数"])
    sheet.append(["通信工程", "物理", 10])
    sheet.append([None, "政治", 5])
    sheet.merge_cells("A2:A3")
    workbook.save(path)
    parsed = parse_file(path)
    assert parsed.tables[0].headers == ["专业", "选科", "计划数"]
    assert parsed.tables[0].rows[0][0] == "通信工程"
    assert parsed.tables[0].rows[1][0] == "通信工程"


def test_multilevel_headers_are_combined(tmp_path):
    path = tmp_path / "multi.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["专业", "招生计划", "招生计划"])
    sheet.append(["名称", "物理类", "历史类"])
    sheet.append(["通信工程", 20, 0])
    workbook.save(path)
    parsed = parse_file(path)
    table = parsed.tables[0]
    assert table.header_levels == 2
    assert table.headers == ["专业 / 名称", "招生计划 / 物理类", "招生计划 / 历史类"]
    assert table.rows == [["通信工程", "20", "0"]]
