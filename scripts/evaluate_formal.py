"""Frozen evaluation: labels are used for scoring, never passed to RAGService."""
import argparse
import csv
import hashlib
import importlib.metadata
import json
import platform
import re
import sqlite3
import subprocess
import time
from contextlib import closing
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path

import _bootstrap  # noqa: F401
from campus_rag.config import ROOT, load_settings


DEFAULT = ROOT / 'data/eval/independent_v1'
MODES = ['bm25', 'dense', 'hybrid', 'rerank']


def digest(path):
    value = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b''):
            value.update(block)
    return value.hexdigest()


def save(path, data):
    temporary = Path(path).with_suffix('.tmp')
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    temporary.replace(path)


def normalized(text):
    text = re.sub(r'\s+', '', text).replace('（', '(').replace('）', ')').replace('：', ':')
    return re.sub(r'[—–－-]+', '-', text)


def raw_text(path):
    if path.suffix == '.pdf':
        import pymupdf
        with pymupdf.open(path) as document:
            return '\n'.join(page.get_text() for page in document)
    if path.suffix in {'.html', '.htm'}:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(path.read_text(encoding='utf-8'), 'lxml')
        return soup.get_text(' ', strip=True)
    raise ValueError('Unsupported raw source: ' + str(path))


def raw_sheet(path):
    """Independent XLSX check, not the project's table parser."""
    import openpyxl
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = workbook.active
        expanded = {}
        for merged in sheet.merged_cells.ranges:
            value = sheet.cell(merged.min_row, merged.min_col).value
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    expanded[row, col] = value
        # This gold source has two header rows, as inspected before model evaluation.
        def cell_text(r, c):
            value = expanded.get((r, c), sheet.cell(r, c).value)
            return '' if value is None else str(value).strip()
        return [[cell_text(r, c) for c in range(1, sheet.max_column + 1)] for r in range(3, sheet.max_row + 1)]
    finally:
        workbook.close()


def prepare(directory):
    if (directory / 'manifest.json').exists():
        raise SystemExit('Already frozen; create a new dataset version instead of overwriting.')
    data = json.loads((directory / 'questions.json').read_text(encoding='utf-8'))
    settings = load_settings()
    gold, catalog, raw_cache = [], {}, {}
    with closing(sqlite3.connect(settings.database_path.resolve().as_uri() + '?mode=ro', uri=True)) as connection:
        connection.row_factory = sqlite3.Row
        for case in data['cases']:
            entry, groups = dict(case), []
            for source in case['sources']:
                doc = dict(connection.execute('SELECT * FROM documents WHERE id=?', (source['doc_id'],)).fetchone())
                path = Path(doc['local_path'])
                if doc['id'] not in catalog:
                    catalog[doc['id']] = {'doc_id': doc['id'], 'title': doc['title'], 'source_url': doc['source_url'],
                        'published_at': doc['published_at'], 'local_path': str(path.relative_to(ROOT)), 'raw_sha256': digest(path)}
                if 'table_id' in source:
                    rows = connection.execute('SELECT id,row_index FROM table_rows WHERE table_id=? ORDER BY row_index', (source['table_id'],)).fetchall()
                    if doc['id'] not in raw_cache:
                        raw_cache[doc['id']] = raw_sheet(path)
                    raw_rows = raw_cache[doc['id']]
                    stored = [[r[0] for r in connection.execute('SELECT raw_value FROM table_cells WHERE row_id=? ORDER BY column_id', (row['id'],))] for row in rows]
                    if len(raw_rows) != len(stored) or any([normalized(v) for v in a] != [normalized(v) for v in b] for a, b in zip(raw_rows, stored)):
                        raise ValueError(case['id'] + ': complete raw XLSX differs from SQLite')
                    selected = [i for i, row in enumerate(rows) if row['id'] in source['row_ids']]
                    if len(selected) != len(source['row_ids']):
                        raise ValueError('Invalid gold row IDs')
                    groups.append({**source, 'raw_verified': True, 'raw_excel_rows': [i + 3 for i in selected],
                        'raw_selected_values': [raw_rows[i] for i in selected], 'full_table_rows': len(raw_rows),
                        'chunk_ids': [r[0] for r in connection.execute('SELECT id FROM chunks WHERE table_id=?', (source['table_id'],))]})
                else:
                    if doc['id'] not in raw_cache:
                        raw_cache[doc['id']] = raw_text(path)
                    if normalized(source['quote']) not in normalized(raw_cache[doc['id']]):
                        raise ValueError(case['id'] + ': quote absent from raw source: ' + source['quote'])
                    chunks = connection.execute("SELECT id,text FROM chunks WHERE doc_id=? AND kind='text'", (doc['id'],)).fetchall()
                    ids = [r['id'] for r in chunks if normalized(source['quote']) in normalized(r['text'])]
                    if not ids:
                        raise ValueError(case['id'] + ': quote absent from frozen chunks')
                    groups.append({**source, 'raw_verified': True, 'chunk_ids': ids})
            entry['evidence_groups'] = groups
            gold.append(entry)
    save(directory / 'gold.json', {'description': data['description'], 'cases': gold})
    save(directory / 'source_catalog.json', list(catalog.values()))
    with (directory / 'questions.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(['问题'])
        writer.writerows([case['question']] for case in gold)
    tracked = [*sorted((ROOT / 'src/campus_rag').rglob('*.py')), ROOT / 'src/campus_rag/database/schema.sql',
        ROOT / 'config/settings.json', Path(__file__), directory / 'questions.json', directory / 'gold.json',
        directory / 'source_catalog.json', directory / 'protocol.md', settings.database_path,
        settings.index_dir / 'records.jsonl', settings.index_dir / 'vectors.npy']
    tracked += [ROOT / source['local_path'] for source in catalog.values()]
    for model in [settings.generator_model, settings.embedding_model, settings.reranker_model]:
        tracked += [p for p in model.iterdir() if p.is_file() and (p.suffix in {'.json', '.safetensors'} or p.name == 'vocab.txt')]
    manifest = {'created_at': datetime.now(timezone.utc).isoformat(),
        'baseline_commit': subprocess.check_output(['git', '-c', 'safe.directory=' + ROOT.as_posix(), 'rev-parse', 'HEAD'], text=True).strip(),
        'files': {str(p.resolve().relative_to(ROOT)): digest(p) for p in tracked}, 'settings': asdict(settings),
        'python': platform.python_version(), 'packages': {p: importlib.metadata.version(p) for p in ['torch', 'transformers', 'sentence-transformers', 'numpy', 'pytest']},
        'case_count': len(gold), 'source_document_count': len(catalog)}
    save(directory / 'manifest.json', manifest)
    print(json.dumps({'frozen_cases': len(gold), 'source_documents': len(catalog)}, ensure_ascii=False))


def verify(directory):
    manifest = json.loads((directory / 'manifest.json').read_text(encoding='utf-8'))
    for name, expected in manifest['files'].items():
        if digest(ROOT / name) != expected:
            raise SystemExit('Frozen file changed: ' + name)
    return manifest


def retrieval_metrics(ids, case):
    groups = [set(g['chunk_ids']) for g in case['evidence_groups']]
    relevant = set().union(*groups) if groups else set()
    result = {}
    for k in [5, 10, 20]:
        found = set(ids[:k])
        result['hit@' + str(k)] = int(bool(found & relevant))
        result['recall@' + str(k)] = len(found & relevant) / len(relevant) if relevant else None
        result['evidence_group_recall@' + str(k)] = sum(bool(found & group) for group in groups) / len(groups) if groups else None
        result['all_evidence_hit@' + str(k)] = int(bool(groups) and all(found & group for group in groups))
    rank = next((i for i, value in enumerate(ids, 1) if value in relevant), None)
    result['rr@20'] = 1 / rank if rank else 0
    return result


class EvaluationIndex:
    def __init__(self, index):
        self.index, self.last = index, {}

    def search(self, query, bm25_top_k, dense_top_k, fusion_top_k, final_top_k, **kwargs):
        self.last = {}
        for mode in MODES:
            started = time.perf_counter()
            found = self.index.search(query, bm25_top_k, dense_top_k, fusion_top_k, max(20, final_top_k), mode=mode, **kwargs)
            self.last[mode] = {'evidence': found, 'seconds': time.perf_counter() - started}
        return self.last['rerank']['evidence'][:final_top_k]


def process_memory():
    if platform.system() != 'Windows':
        return None
    import ctypes
    from ctypes import wintypes
    class Counters(ctypes.Structure):
        _fields_ = [('cb', wintypes.DWORD), ('faults', wintypes.DWORD)] + [(name, ctypes.c_size_t) for name in ['peak_rss', 'rss', 'peak_paged', 'paged', 'peak_nonpaged', 'nonpaged', 'pagefile', 'peak_pagefile']]
    counters = Counters()
    counters.cb = ctypes.sizeof(counters)
    ctypes.windll.kernel32.GetCurrentProcess.restype = wintypes.HANDLE
    ctypes.windll.psapi.GetProcessMemoryInfo.argtypes = [wintypes.HANDLE, ctypes.c_void_p, wintypes.DWORD]
    if ctypes.windll.psapi.GetProcessMemoryInfo(ctypes.windll.kernel32.GetCurrentProcess(), ctypes.byref(counters), counters.cb):
        return {'working_set_bytes': counters.rss, 'process_peak_working_set_bytes': counters.peak_rss}
    return None


def run(directory, resume=False):
    manifest = verify(directory)
    output = directory / 'run1'
    output.mkdir(exist_ok=True)
    result_path = output / 'results.json'
    if result_path.exists() and not resume:
        raise SystemExit('Results exist; --resume only completes missing cases.')
    import socket
    attempts = []
    def denied(*args, **kwargs):
        attempts.append('blocked')
        raise RuntimeError('Network disabled during formal evaluation')
    socket.socket.connect = denied
    socket.socket.connect_ex = denied
    import numpy as np
    import torch
    from campus_rag.index.hybrid import HybridIndex
    from campus_rag.rag.generator import TransformersGenerator
    from campus_rag.rag.service import RAGService
    torch.manual_seed(0)
    np.random.seed(0)
    settings = load_settings()
    started = time.perf_counter()
    generator = TransformersGenerator(settings.generator_model)
    index = EvaluationIndex(HybridIndex.load(settings.index_dir, settings.embedding_model, settings.reranker_model))
    loaded = time.perf_counter() - started
    cases = json.loads((directory / 'gold.json').read_text(encoding='utf-8'))['cases']
    report = json.loads(result_path.read_text(encoding='utf-8')) if result_path.exists() else {
        'baseline_commit': manifest['baseline_commit'], 'manifest_sha256': digest(directory / 'manifest.json'),
        'started_at': datetime.now(timezone.utc).isoformat(), 'model_init_seconds': loaded, 'results': [],
        'network_policy': 'process-local socket connections blocked; not a physical disconnect test'}
    if report['manifest_sha256'] != digest(directory / 'manifest.json'):
        raise SystemExit('Resume report does not belong to the frozen manifest')
    report['device'] = torch.cuda.get_device_name() if torch.cuda.is_available() else 'CPU'
    done = {r['id'] for r in report['results']}
    with closing(sqlite3.connect(settings.database_path.resolve().as_uri() + '?mode=ro', uri=True)) as connection:
        service = RAGService(settings, index, generator, connection)
        for case in cases:
            if case['id'] in done:
                continue
            started = time.perf_counter()
            index.last = {}
            service.planner.last_raw = ''
            if torch.cuda.is_available():
                torch.cuda.synchronize()
                torch.cuda.reset_peak_memory_stats()
            row = {'id': case['id'], 'category': case['category'], 'question': case['question']}
            try:
                # No gold answer, expected value, or relevant ID is passed to the system.
                result = service.answer(case['question'])
                row.update(result=asdict(result), plan=asdict(service.last_plan), raw_plan=getattr(service.planner, 'last_raw', ''))
            except Exception as exc:
                row['error'] = type(exc).__name__ + ': ' + str(exc)
            if torch.cuda.is_available():
                torch.cuda.synchronize()
                row['gpu_peak_allocated_bytes'] = torch.cuda.max_memory_allocated()
                row['gpu_peak_reserved_bytes'] = torch.cuda.max_memory_reserved()
            row['seconds_with_retrieval_comparisons'] = round(time.perf_counter() - started, 3)
            row['retrieval'] = {mode: {'seconds': value['seconds'], 'ids': [e.evidence_id for e in value['evidence']],
                'doc_ids': [e.doc_id for e in value['evidence']], 'metrics': retrieval_metrics([e.evidence_id for e in value['evidence']], case) if not case.get('refuse') else None}
                for mode, value in index.last.items()}
            extra = sum(value['seconds'] for key, value in index.last.items() if key != 'rerank')
            row['service_seconds_estimate'] = round(row['seconds_with_retrieval_comparisons'] - extra, 3)
            row['memory'] = process_memory()
            report['results'].append(row)
            report['network_attempt_count'] = len(attempts)
            save(result_path, report)
            with (output / 'results.csv').open('w', encoding='utf-8-sig', newline='') as handle:
                writer = csv.DictWriter(handle, fieldnames=['问题', '答案', '参考资料'])
                writer.writeheader()
                for item in report['results']:
                    answer = item.get('result', {})
                    writer.writerow({'问题': item['question'], '答案': answer.get('answer', '运行异常：' + item.get('error', '')), '参考资料': '\n'.join(answer.get('references', []))})
            print(json.dumps({'id': row['id'], 'seconds': row['service_seconds_estimate'], 'error': row.get('error')}, ensure_ascii=False), flush=True)
    report['finished_at'] = datetime.now(timezone.utc).isoformat()
    save(result_path, report)
    verify(directory)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--directory', type=Path, default=DEFAULT)
    parser.add_argument('--prepare', action='store_true')
    parser.add_argument('--run', action='store_true')
    parser.add_argument('--resume', action='store_true')
    args = parser.parse_args()
    if args.prepare == args.run:
        parser.error('choose exactly one of --prepare or --run')
    prepare(args.directory) if args.prepare else run(args.directory, args.resume)


if __name__ == '__main__':
    main()
